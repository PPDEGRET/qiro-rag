from pathlib import Path

from qiro_rag.index import EvidenceIndex
from qiro_rag.ingest import ingest_path


def test_ingests_docx_and_xlsx(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import Workbook

    docs = tmp_path / "docs"
    pack = tmp_path / "pack"
    docs.mkdir()

    document = Document()
    document.add_paragraph("Product specification: ReBox mailer contains 82% recycled content.")
    document.save(docs / "spec.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet.append(["sku", "claim", "evidence"])
    sheet.append(["RB-1", "recyclable", "Material is recyclable where PE film collection exists."])
    workbook.save(docs / "evidence.xlsx")

    summary = ingest_path(docs, pack, reset=True)
    hits = retrieve_all_text(pack)

    assert summary.indexed_documents == 2
    assert "82% recycled content" in hits
    assert "PE film collection" in hits


def retrieve_all_text(pack: Path) -> str:
    index = EvidenceIndex(pack)
    return "\n".join(hit.text for hit in index.all_chunks())
